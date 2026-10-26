"""Concise, source-checked processing and reproduction notes for Figure 1."""
from docx.shared import Inches, Pt
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def link(p, label, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
    r.append(props); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r)
    p._p.append(h)


def add_methods(d):
    section = d.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name, size, before, after in [('Methods body',10.5,0,6), ('Methods heading',12,10,5), ('Methods title',16,0,10), ('Methods code',9,0,5)]:
        s = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = d.styles['Normal']
        s.font.name = 'Consolas' if name == 'Methods code' else 'Calibri'
        s.font.size = Pt(size); s.font.bold = name in ['Methods heading','Methods title']
        f=s.paragraph_format; f.space_before=Pt(before); f.space_after=Pt(after); f.line_spacing=1.08
        f.keep_with_next = name in ['Methods heading','Methods title']
        f.widow_control = True
        if name in ['Methods heading','Methods title']:
            level=OxmlElement('w:outlineLvl'); level.set(qn('w:val'),'0' if name=='Methods title' else '1')
            f._element.append(level)
    def p(text): return d.add_paragraph(text, style='Methods body')
    def heading(text): d.add_paragraph(text,style='Methods heading')

    d.add_paragraph('Data processing and reproducibility',style='Methods title')
    heading('ChIP-seq sources')
    q=p('The 48-hour EOMES ChIP-seq data are from HUES8 human embryonic stem cells differentiating toward definitive endoderm: ')
    link(q,'Luo et al., Nature Genetics (2023)', 'https://doi.org/10.1038/s41588-023-01450-7')
    q.add_run(', GEO '); link(q,'GSE213394','https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE213394')
    q.add_run('. Replicate 1 is GSM6585561 / '); link(q,'SRR21586161','https://www.ebi.ac.uk/ena/browser/view/SRR21586161')
    q.add_run('; replicate 2 is GSM6585562 / '); link(q,'SRR21586115','https://www.ebi.ac.uk/ena/browser/view/SRR21586115')
    q.add_run(' (paired-end, 101 bp). Raw reads are archived in SRA and available through these ENA run pages.')
    q=p('The selected external input is GSM3381961 / ')
    link(q,'SRR7814534','https://www.ebi.ac.uk/ena/browser/view/SRR7814534')
    q.add_run(' (single-end, 51 bp; control DE day 1), from ')
    link(q,'Li et al., Nature Genetics (2019)','https://doi.org/10.1038/s41588-019-0408-9')
    q.add_run('. Its '); link(q,'ENA FASTQ download','https://ftp.sra.ebi.ac.uk/vol1/fastq/SRR781/004/SRR7814534/SRR7814534.fastq.gz')
    q.add_run(' was MD5-verified. It was selected in the earlier background-coverage comparison and is not an experimentally matched input. Only the selected 48-hour result is used here.')

    heading('Read processing and reproducible peaks')
    p('Reads were aligned to GRCh38 no-alt with Bowtie2 2.5.5 (default alignment settings). SAMtools 1.24 name-sorted reads, repaired mate information, coordinate-sorted and removed duplicates. Retained alignments had MAPQ ≥30, were on chr1–22/X/Y and excluded hg38 blacklist v2 overlaps. ChIP BAM headers record proper-pair filtering (-f 2 -F 1804); the single-end input used -F 3844. The reference FASTA/index and annotation are recorded in the source provenance.')
    p('MACS2 2.2.9.1 called peaks separately for each replicate and pooled ChIP against the selected input, using -f BAM --nomodel --extsize 200 --shift 0 --keep-dup all -g hs -p 0.001. Thus, paired-end ChIP was treated as read-1 tags with fixed 200-bp extension, not BAMPE fragments. After blacklist exclusion, IDR 2.0.4.2 compared biological replicates ranked by peak p-value, with pooled peaks as the oracle. Global IDR ≤0.05 and pooled peak q ≤0.05 retained 43,700 peaks.')

    heading('EOMES target definition')
    p('Using GENCODE v44 basic protein-coding and lncRNA genes, each retained pooled peak summit was linked to every gene-level TSS within 50 kb. The gene score is the sum of peak fold enrichment × (1 − |summit − TSS|/50,000). This is a distance-weighted sum, not an average. Repeated symbols take the maximum locus score; score ties are alphabetical. The first 500 and 1,500 unique uppercase symbols define the fixed target sets. Missing or expression-ineligible genes are not replaced. These are candidate binding targets based on proximity.')

    heading('MorPhiC DE inputs')
    p('Saved MorPhiC gene-level DE exports are reused without refitting. The upstream notebook used TMM normalization, limma-voom group models, empirical-Bayes contrasts and clone duplicateCorrelation where applicable. q75_pass requires a positive 75th percentile of counts across the original sample matrix. Eligibility additionally requires maximum group-mean TMM CPM ≥1; DE requires Storey q <0.05 and |log2FC| >1. LogFC is perturbation/control. Degron compares IAA-treated tagged clones with their matched DMSO controls; CRISPRi compares doxycycline-induced with uninduced cells. The two KO methods share WT controls.')

    d.add_page_break()
    d.add_paragraph('Figure calculations and rerunning',style='Methods title')
    heading('Panel A: directional target enrichment')
    p('Within each contrast, the background is all expression-eligible RNA symbols, matched exactly in uppercase; missing symbols retain Ensembl IDs. A symbol is directionally DE if any eligible locus qualifies (opposite-sign loci may place it in both sets). For each target size and DE direction, a one-sided Fisher exact test compares target membership with DE status. Fold enrichment = observed overlap / (eligible target count × directional DE count / background count). Stars indicate BH q <0.05. The original correction over 20 tests (five contrasts × two sizes × two directions) is retained after hiding the parental-control degron column.')

    heading('Panels B–C: cell-program enrichment')
    q=p('The fixed library contains 96 signatures: 21 human gastrula signatures using the first 50 published markers from supplementary tables 1, 8, 11 and 13 of ')
    link(q,'Tyser et al., Nature (2021)','https://doi.org/10.1038/s41586-021-04158-y')
    q.add_run('; and 75 main fetal cell-type signatures from ')
    link(q,'Cao et al., Science (2020)','https://doi.org/10.1126/science.aba7721')
    q.add_run(', obtained from MSigDB C8 v2025.1.Hs. The frozen GMT and source metadata are in inputs/reference/. The figure displays 16 signatures; testing includes all eligible signatures.')
    from update_wt_documents import METHODS
    for text in METHODS:
        p(text)

    heading('Reproduce the released figure')
    p('From the repository root, install requirements.txt plus the R packages in R_packages.txt, then run:')
    d.add_paragraph('python3 run.py',style='Methods code')
    p('Use --rscript /path/to/Rscript if needed. To regenerate only this four-column display and document from existing tables:')
    d.add_paragraph('python3 scripts/render_results.py --single-degron\npython3 scripts/build_document.py --single-degron',style='Methods code')
    p('The release is offline and starts from frozen final ChIP peaks, links, scores, target rankings and original DE tables. It does not rerun raw-read alignment, input selection or observed RNA DE models. WT-versus-WT models use the frozen inputs in inputs/wt_reference/. Upstream code remains in chipseq_eomes/ (align_fastq.py, pipeline.py); historical commands are in logs/commands.jsonl and BAM headers. Raw-data paths and reference checksums are in provenance/. The recorded ChIP filter above describes the historical BAMs; the optional realignment script additionally excludes supplementary alignments.')
    p('Numerical outputs are in results/tables/; the figure and document use the suffix _single_degron. config.json records thresholds and displayed signatures. Input and output checksums, validation and software versions are in provenance/, release_manifest.json and qc/. Tested core versions: Python 3.9.6, NumPy 1.26.3, pandas 2.2.0, SciPy 1.12.0, Matplotlib 3.8.2; R 4.6.1, fgsea 1.38.0; python-docx 1.2.0. ReMap is retained only as a separate sensitivity analysis.')

