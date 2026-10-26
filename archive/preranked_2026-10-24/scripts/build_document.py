"""Build the Word figure/caption document and concise methods for the single-Degron version."""
import json
import argparse
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from common import ROOT

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--single-degron',action='store_true')
    args=parser.parse_args()
    suffix='_single_degron' if args.single_degron else ''
    figure_width=7.9 if args.single_degron else 9.1
    caption=json.loads((ROOT/f'results/figures/captions{suffix}.json').read_text())['primary']
    d=Document();section=d.sections[0]
    section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Inches(11);section.page_height=Inches(8.5)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(.5)
    section.header_distance=section.footer_distance=Inches(.3)
    def style(name,size,before,after,line=1.1,color='000000',bold=False):
        st=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color);st.font.bold=bold
        pf=st.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line
        pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
        return st
    style('Normal',11,0,6)
    style('Title',24,0,10);style('Subtitle',12,0,6)
    style('Heading 1',16,16,8,color='2E74B5',bold=True)
    style('Heading 2',13,12,6,color='2E74B5',bold=True)
    style('Heading 3',12,8,4,color='1F4D78',bold=True)
    style('Header',9,0,0);style('Footer',9,0,0)
    st=style('Figure image',1,0,0,line=1.0);st.paragraph_format.keep_with_next=True
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    st=style('Caption',9.5,8,0);st.paragraph_format.keep_together=True
    p=d.add_paragraph(style='Figure image')
    shape=p.add_run().add_picture(str(ROOT/f'results/figures/figure_1_chip48_and_cell_fate{suffix}.png'),width=Inches(figure_width))
    shape._inline.docPr.set('descr','Panel A: separate upregulated and downregulated DE enrichment for the 48-hour top 500 and top 1,500 ChIP targets. Panels B and C: cell-fate GSEA before and after removing the 48-hour top 1,500 targets. Columns are perturbation methods.')
    p=d.add_paragraph(style='Caption')
    p.add_run('Figure 1. '+caption['title']+' ').bold=True
    p.add_run(caption['text'])
    if args.single_degron:
        from document_methods import add_methods
        add_methods(d)
    d.core_properties.title='48-hour EOMES targets and cell-fate programs'
    d.core_properties.author='';d.core_properties.last_modified_by=''
    out=ROOT/f'results/documents/figure_1_with_caption{suffix}.docx'
    d.save(out)
    design=dict(preset='standard_business_brief',named_overrides=dict(
        scientific_figure_page=dict(size='US Letter landscape',width_in=11,height_in=8.5,margins_in=.5,header_footer_in=.3),
        scientific_caption=dict(font='Calibri',size_pt=9.5,before_pt=8,after_pt=0,line_spacing=1.1),
        scientific_figure=dict(width_in=figure_width,alignment='center',paragraph_font_pt=1,spacing_pt=0)))
    if args.single_degron:
        design['named_overrides']['methods_pages']=dict(size='US Letter portrait',margins_lr_in=.75,margins_tb_in=.7,font='Calibri',body_pt=10.5,line_spacing=1.08,paragraph_after_pt=6,heading_pt=12,title_pt=16,code_font='Consolas',code_pt=9)
    (ROOT/f'qc/document_design{suffix}.json').write_text(json.dumps(design,indent=2)+'\n')
    print(out)

if __name__=='__main__':main()
